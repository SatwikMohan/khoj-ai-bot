import argparse
import csv
import hashlib
import io
import json
import os
import time
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from pathlib import Path

import chromadb
import docx2txt
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pptx import Presentation

from services.embedding_service import PromptedOllamaEmbeddings, embedding_profile
from services.lexical_service import add_lexical_chunks, delete_lexical_ids, reset_lexical_collection


PROJECT_ROOT = Path(__file__).resolve().parent
REPO_ROOT = PROJECT_ROOT
RAW_DATA_DIR = REPO_ROOT / "raw_data_files"
VECTOR_DB_DIR = REPO_ROOT / "vector_db"
MANIFEST_FILE_NAME = "ingestion_manifest.json"
INGESTION_VERSION = 6
OCR_WARNING_KEYS: set[str] = set()
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".css",
    ".html",
    ".htm",
    ".xml",
    ".yaml",
    ".yml",
}
IMAGE_EXTENSIONS = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}


def load_environment() -> None:
    load_dotenv(PROJECT_ROOT / ".env")


def resolve_path(value: str | None, default: Path) -> Path:
    if not value:
        return default.resolve()

    path = Path(value)
    if path.is_absolute():
        return path

    for base_dir in (REPO_ROOT, PROJECT_ROOT):
        candidate = (base_dir / path).resolve()
        if candidate.exists():
            return candidate

    return (REPO_ROOT / path).resolve()


def chunk_settings() -> dict:
    profile = embedding_profile()
    return {
        "version": INGESTION_VERSION,
        "embedding_provider": "ollama",
        "embedding_model": profile.model,
        "embedding_profile": profile.as_dict(),
        "ollama_base_url": os.getenv("OLLAMA_BASE_URL", "http://localhost:11434"),
        "chunk_size": int(os.getenv("CHUNK_SIZE", "850")),
        "chunk_overlap": int(os.getenv("CHUNK_OVERLAP", "150")),
        "collection_name": os.getenv("CHROMA_COLLECTION_NAME", "texmin_qa"),
        "ocr_enabled": ocr_enabled(),
        "ocr_lang": os.getenv("OCR_LANG", "eng+hin"),
        "ocr_pdf_dpi": int(os.getenv("OCR_PDF_DPI", "200")),
    }


def relative_source(path: Path) -> str:
    try:
        return path.relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return path.relative_to(PROJECT_ROOT).as_posix()


def raw_relative_source(path: Path, raw_data_dir: Path = RAW_DATA_DIR) -> str:
    return path.relative_to(raw_data_dir).as_posix()


def base_metadata(path: Path, raw_data_dir: Path = RAW_DATA_DIR) -> dict:
    raw_relative_path = raw_relative_source(path, raw_data_dir)
    folder_path = Path(raw_relative_path).parent.as_posix()
    if folder_path == ".":
        folder_path = ""

    return {
        "source": relative_source(path),
        "raw_source": raw_relative_path,
        "folder_path": folder_path,
        "folder_depth": 0 if not folder_path else len(Path(folder_path).parts),
        "file_name": path.name,
        "file_path": str(path),
        "extension": path.suffix.lower(),
    }


def read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="ignore")


def env_flag(name: str, default: bool = True) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() not in {"0", "false", "no", "off"}


def ocr_enabled() -> bool:
    return env_flag("OCR_ENABLED", True)


def warn_once(key: str, message: str) -> None:
    if key in OCR_WARNING_KEYS:
        return
    OCR_WARNING_KEYS.add(key)
    print(message)


def normalized_text_length(text: str) -> int:
    return len(" ".join((text or "").split()))


def ocr_image(image, source_label: str = "image") -> str:
    if not ocr_enabled():
        return ""

    try:
        import pytesseract
        from PIL import ImageOps
    except ImportError as exc:
        warn_once(
            "ocr-python-deps",
            f"OCR skipped for {source_label}: install pytesseract and Pillow. {exc}",
        )
        return ""

    prepared_image = ImageOps.grayscale(image)
    configured_lang = os.getenv("OCR_LANG", "eng+hin").strip() or "eng"
    languages = [configured_lang]
    if configured_lang != "eng":
        languages.append("eng")

    config = os.getenv("OCR_TESSERACT_CONFIG", "--psm 6")
    for index, language in enumerate(languages):
        try:
            return pytesseract.image_to_string(prepared_image, lang=language, config=config).strip()
        except pytesseract.pytesseract.TesseractNotFoundError as exc:
            warn_once(
                "ocr-tesseract-missing",
                f"OCR skipped for {source_label}: Tesseract is not installed or not on PATH. {exc}",
            )
            return ""
        except pytesseract.pytesseract.TesseractError as exc:
            if index + 1 < len(languages):
                warn_once(
                    f"ocr-lang-{language}",
                    f"OCR language '{language}' failed for {source_label}; falling back to English. {exc}",
                )
                continue
            warn_once(
                f"ocr-error-{language}",
                f"OCR failed for {source_label} with language '{language}'. {exc}",
            )
            return ""

    return ""


def load_image(path: Path, raw_data_dir: Path) -> list[Document]:
    try:
        from PIL import Image
    except ImportError as exc:
        warn_once("ocr-pillow-missing", f"OCR skipped for {relative_source(path)}: install Pillow. {exc}")
        return []

    try:
        with Image.open(path) as image:
            text = ocr_image(image, relative_source(path))
    except Exception as exc:
        print(f"Could not OCR image {relative_source(path)}: {exc}")
        return []

    if not text.strip():
        print(f"Skipping image with no OCR text: {relative_source(path)}")
        return []

    metadata = base_metadata(path, raw_data_dir)
    metadata.update({"ocr": True, "ocr_source": "image"})
    return [Document(page_content=text, metadata=metadata)]


def load_pdf_ocr_pages(path: Path, raw_data_dir: Path, page_numbers: list[int] | None) -> list[Document]:
    if page_numbers == [] or not ocr_enabled():
        return []

    try:
        import fitz
        from PIL import Image
    except ImportError as exc:
        warn_once(
            "ocr-pdf-deps",
            f"PDF OCR skipped for {relative_source(path)}: install PyMuPDF and Pillow. {exc}",
        )
        return []

    documents = []
    dpi = int(os.getenv("OCR_PDF_DPI", "200"))
    try:
        with fitz.open(str(path)) as pdf:
            if page_numbers is None:
                page_numbers = list(range(pdf.page_count))
            for page_number in page_numbers:
                if page_number < 0 or page_number >= pdf.page_count:
                    continue

                page = pdf.load_page(page_number)
                pixmap = page.get_pixmap(dpi=dpi, alpha=False)
                with Image.open(io.BytesIO(pixmap.tobytes("png"))) as image:
                    text = ocr_image(image, f"{relative_source(path)} page {page_number + 1}")

                if not text.strip():
                    continue

                metadata = base_metadata(path, raw_data_dir)
                metadata.update(
                    {
                        "page": page_number,
                        "ocr": True,
                        "ocr_source": "pdf_page",
                    }
                )
                documents.append(Document(page_content=text, metadata=metadata))
    except Exception as exc:
        print(f"Could not OCR PDF {relative_source(path)}: {exc}")

    return documents


def file_hash(path: Path) -> str:
    digest = hashlib.sha1()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_csv(path: Path) -> str:
    rows = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        for row in reader:
            rows.append(", ".join(cell.strip() for cell in row if cell is not None))
    return "\n".join(rows)


def load_json(path: Path) -> str:
    data = json.loads(read_text(path))
    return json.dumps(data, indent=2, ensure_ascii=False)


def load_docx(path: Path) -> str:
    return docx2txt.process(str(path)) or ""


def load_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                parts.append("\t".join(values))
    return "\n".join(parts)


def load_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    parts = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts.append(f"Slide {slide_number}")
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                parts.append(text)
    return "\n".join(parts)


def load_pdf(path: Path, raw_data_dir: Path) -> list[Document]:
    documents = []
    min_text_chars = int(os.getenv("MIN_EXTRACTED_TEXT_CHARS", "80"))
    try:
        loaded_documents = PyPDFLoader(str(path)).load()
    except Exception as exc:
        print(f"Could not extract PDF text from {relative_source(path)}; trying OCR. {exc}")
        return load_pdf_ocr_pages(path, raw_data_dir, None)

    low_text_pages = []
    for fallback_page_number, document in enumerate(loaded_documents):
        try:
            page_number = int(document.metadata.get("page", fallback_page_number))
        except (TypeError, ValueError):
            page_number = fallback_page_number
        if normalized_text_length(document.page_content) < min_text_chars:
            low_text_pages.append(page_number)
            continue
        document.metadata.update(base_metadata(path, raw_data_dir))
        documents.append(document)

    documents.extend(load_pdf_ocr_pages(path, raw_data_dir, low_text_pages))
    return documents


def load_file(path: Path, raw_data_dir: Path) -> list[Document]:
    suffix = path.suffix.lower()
    metadata = base_metadata(path, raw_data_dir)

    if suffix == ".pdf":
        return load_pdf(path, raw_data_dir)
    if suffix in IMAGE_EXTENSIONS:
        return load_image(path, raw_data_dir)
    if suffix == ".csv":
        text = load_csv(path)
    elif suffix == ".json":
        text = load_json(path)
    elif suffix == ".docx":
        text = load_docx(path)
    elif suffix in {".xlsx", ".xlsm"}:
        text = load_xlsx(path)
    elif suffix == ".pptx":
        text = load_pptx(path)
    elif suffix in TEXT_EXTENSIONS:
        text = read_text(path)
    else:
        print(f"Skipping unsupported file type: {relative_source(path)}")
        return []

    if not text.strip():
        print(f"Skipping empty file: {relative_source(path)}")
        return []

    return [Document(page_content=text, metadata=metadata)]


def iter_raw_files(raw_data_dir: Path) -> list[Path]:
    raw_data_dir = raw_data_dir.resolve()
    if not raw_data_dir.exists():
        raise RuntimeError(f"Raw data folder does not exist: {raw_data_dir}")

    return sorted(
        path
        for path in raw_data_dir.rglob("*")
        if path.is_file()
        and path.resolve().is_relative_to(raw_data_dir)
        and not any(part.startswith(".") for part in path.relative_to(raw_data_dir).parts)
    )


def load_documents(raw_data_dir: Path) -> list[Document]:
    raw_data_dir = raw_data_dir.resolve()
    documents = []
    for path in iter_raw_files(raw_data_dir):
        try:
            documents.extend(load_file(path, raw_data_dir))
        except Exception as exc:
            print(f"Could not load {relative_source(path)}: {exc}")
    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    settings = chunk_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings["chunk_size"],
        chunk_overlap=settings["chunk_overlap"],
        separators=[
            "\n\n",
            "\n",
            ". ",
            "? ",
            "! ",
            "; ",
            ", ",
            " ",
            "",
        ],
        add_start_index=True,
    )
    chunks = splitter.split_documents(documents)
    for index, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = index
    return chunks


def chunk_id(chunk: Document) -> str:
    identity = {
        "source": chunk.metadata.get("source"),
        "page": chunk.metadata.get("page"),
        "start_index": chunk.metadata.get("start_index"),
        "content": chunk.page_content,
    }
    return hashlib.sha1(json.dumps(identity, sort_keys=True).encode("utf-8")).hexdigest()


def manifest_path(vector_db_dir: Path) -> Path:
    return vector_db_dir / MANIFEST_FILE_NAME


def load_manifest(vector_db_dir: Path) -> dict:
    path = manifest_path(vector_db_dir)
    if not path.exists():
        return {"settings": {}, "files": {}}

    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {"settings": {}, "files": {}}


def save_manifest(vector_db_dir: Path, manifest: dict) -> None:
    path = manifest_path(vector_db_dir)
    temporary_path = path.with_suffix(".tmp")
    temporary_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    temporary_path.replace(path)


def create_vector_store(
    collection_name: str | None = None,
    reset_collection: bool = False,
) -> tuple[Chroma, Path, str]:
    collection_name = collection_name or os.getenv("CHROMA_COLLECTION_NAME", "texmin_qa")
    vector_db_dir = resolve_path(os.getenv("VECTOR_DB_DIR"), VECTOR_DB_DIR)

    vector_db_dir.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(vector_db_dir))
    if reset_collection:
        try:
            client.delete_collection(collection_name)
        except Exception:
            pass

    profile = embedding_profile()
    embeddings = PromptedOllamaEmbeddings(
        profile=profile,
        base_url=os.getenv("OLLAMA_BASE_URL", "http://localhost:11434"),
        keep_alive=int(os.getenv("OLLAMA_KEEP_ALIVE", "1800")),
    )
    vector_store = Chroma(
        client=client,
        collection_name=collection_name,
        embedding_function=embeddings,
        collection_metadata={
            "hnsw:space": "cosine",
            "embedding_model": profile.model,
            "embedding_profile": profile.fingerprint,
            "ingestion_version": INGESTION_VERSION,
        },
    )
    return vector_store, vector_db_dir, collection_name


def delete_ids(vector_store: Chroma, ids: list[str]) -> None:
    if ids:
        vector_store.delete(ids=ids)


def add_chunks(
    vector_store: Chroma,
    vector_db_dir: Path,
    collection_name: str,
    chunks: list[Document],
) -> list[str]:
    batch_size = int(os.getenv("EMBEDDING_BATCH_SIZE", "256"))
    ids = [chunk_id(chunk) for chunk in chunks]

    for start in range(0, len(chunks), batch_size):
        end = start + batch_size
        vector_store.add_documents(chunks[start:end], ids=ids[start:end])
        print(f"Embedded chunks {start + 1}-{min(end, len(chunks))} of {len(chunks)}")
    add_lexical_chunks(vector_db_dir, collection_name, chunks, ids)
    return ids


def versioned_collection_name(base_name: str, settings: dict, force_rebuild: bool) -> str:
    settings_fingerprint = hashlib.sha1(
        json.dumps(settings, sort_keys=True).encode("utf-8")
    ).hexdigest()[:12]
    name = f"{base_name}__{settings_fingerprint}__v{INGESTION_VERSION}"
    if force_rebuild:
        name = f"{name}__{int(time.time())}"
    return name[:63]


def prepare_file(
    path: Path,
    raw_data_dir: Path,
) -> tuple[list[Document], list[Document], str | None]:
    try:
        documents = load_file(path, raw_data_dir)
        return documents, split_documents(documents) if documents else [], None
    except Exception as exc:
        return [], [], str(exc)


def iter_prepared_files(work_items: list[tuple], raw_data_dir: Path):
    workers = max(1, min(8, int(os.getenv("INGESTION_WORKERS", "8"))))
    if workers == 1:
        for item in work_items:
            yield item, prepare_file(item[1], raw_data_dir)
        return

    iterator = iter(work_items)
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="texmin-ingest") as executor:
        in_flight = {}
        for _ in range(workers * 2):
            try:
                item = next(iterator)
            except StopIteration:
                break
            in_flight[executor.submit(prepare_file, item[1], raw_data_dir)] = item

        while in_flight:
            completed, _ = wait(in_flight, return_when=FIRST_COMPLETED)
            for future in completed:
                item = in_flight.pop(future)
                yield item, future.result()
                try:
                    next_item = next(iterator)
                except StopIteration:
                    continue
                in_flight[executor.submit(prepare_file, next_item[1], raw_data_dir)] = next_item


def train(raw_data_dir: Path, force_rebuild: bool = False) -> None:
    load_environment()
    raw_data_dir = raw_data_dir.resolve()

    raw_files = iter_raw_files(raw_data_dir)
    if not raw_files:
        raise RuntimeError(f"No supported documents found in {raw_data_dir}")

    current_settings = chunk_settings()
    vector_db_dir = resolve_path(os.getenv("VECTOR_DB_DIR"), VECTOR_DB_DIR)
    vector_db_dir.mkdir(parents=True, exist_ok=True)
    manifest = load_manifest(vector_db_dir)
    settings_changed = manifest.get("settings") != current_settings
    base_collection_name = os.getenv("CHROMA_COLLECTION_NAME", "texmin_qa")
    building_new_collection = force_rebuild or settings_changed

    if building_new_collection:
        collection_name = versioned_collection_name(
            base_collection_name,
            current_settings,
            force_rebuild,
        )
        print(
            "Building a new collection before switching traffic: "
            f"{collection_name}. The current collection remains available."
        )
        vector_store, vector_db_dir, _ = create_vector_store(
            collection_name=collection_name,
            reset_collection=True,
        )
        reset_lexical_collection(vector_db_dir, collection_name)
        manifest = {
            "schema_version": 2,
            "settings": current_settings,
            "active_collection": collection_name,
            "files": {},
        }
    else:
        collection_name = manifest.get("active_collection") or base_collection_name
        vector_store, vector_db_dir, _ = create_vector_store(collection_name=collection_name)

    manifest.setdefault("files", {})
    manifest.setdefault("failures", {})
    current_files = {raw_relative_source(path, raw_data_dir): path for path in raw_files}

    stale_sources = sorted(set(manifest["files"]) - set(current_files))
    for source in stale_sources:
        delete_ids(vector_store, manifest["files"][source].get("ids", []))
        delete_lexical_ids(
            vector_db_dir,
            collection_name,
            manifest["files"][source].get("ids", []),
        )
        del manifest["files"][source]
        print(f"Removed stale embeddings: {source}")

    changed_count = 0
    skipped_count = 0
    document_count = 0
    chunk_count = 0
    failed_count = 0
    work_items = []

    for source, path in sorted(current_files.items()):
        previous_entry = manifest["files"].get(source)
        stat = path.stat()
        if (
            previous_entry
            and env_flag("FAST_FILE_CHECK", True)
            and previous_entry.get("size") == stat.st_size
            and previous_entry.get("mtime_ns") == stat.st_mtime_ns
        ):
            skipped_count += 1
            continue

        current_hash = file_hash(path)
        if previous_entry and previous_entry.get("hash") == current_hash:
            previous_entry["size"] = stat.st_size
            previous_entry["mtime_ns"] = stat.st_mtime_ns
            skipped_count += 1
            continue

        work_items.append((source, path, current_hash, stat.st_size, stat.st_mtime_ns))

    print(
        f"Preparing {len(work_items)} changed files with "
        f"{max(1, min(8, int(os.getenv('INGESTION_WORKERS', '8'))))} extraction workers."
    )
    for item, prepared in iter_prepared_files(work_items, raw_data_dir):
        source, path, current_hash, file_size, mtime_ns = item
        documents, chunks, preparation_error = prepared
        previous_entry = manifest["files"].get(source)

        if preparation_error:
            failed_count += 1
            manifest["failures"][source] = preparation_error
            print(f"Could not prepare {source}: {preparation_error}")
            continue
        manifest["failures"].pop(source, None)

        if previous_entry:
            delete_ids(vector_store, previous_entry.get("ids", []))
            delete_lexical_ids(
                vector_db_dir,
                collection_name,
                previous_entry.get("ids", []),
            )

        if not documents:
            manifest["files"].pop(source, None)
            continue

        if not chunks:
            manifest["files"].pop(source, None)
            continue

        ids = add_chunks(vector_store, vector_db_dir, collection_name, chunks)
        manifest["files"][source] = {
            "hash": current_hash,
            "size": file_size,
            "mtime_ns": mtime_ns,
            "ids": ids,
            "document_count": len(documents),
            "chunk_count": len(chunks),
        }
        changed_count += 1
        document_count += len(documents)
        chunk_count += len(chunks)
        print(f"Updated embeddings: {source} ({len(chunks)} chunks)")
        if not building_new_collection:
            manifest["settings"] = current_settings
            manifest["schema_version"] = 2
            manifest["active_collection"] = collection_name
            save_manifest(vector_db_dir, manifest)

    if building_new_collection and failed_count:
        raise RuntimeError(
            f"The candidate index has {failed_count} failed source files and was not promoted. "
            "Fix the reported extraction errors and rerun ingestion."
        )

    manifest["settings"] = current_settings
    manifest["schema_version"] = 2
    manifest["active_collection"] = collection_name
    save_manifest(vector_db_dir, manifest)
    print(
        "Training complete. "
        f"Updated {changed_count} files, skipped {skipped_count} unchanged files, "
        f"failed {failed_count} files, removed {len(stale_sources)} stale files, "
        f"embedded {document_count} documents "
        f"into {chunk_count} chunks."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build or refresh the local Chroma vector database from raw_data_files."
    )
    parser.add_argument(
        "--raw-data-dir",
        type=Path,
        default=resolve_path(os.getenv("RAW_DATA_DIR"), RAW_DATA_DIR),
        help="Folder containing source files to embed.",
    )
    parser.add_argument(
        "--force-rebuild",
        action="store_true",
        help="Build a fresh versioned collection and promote it atomically.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    train(args.raw_data_dir.resolve(), force_rebuild=args.force_rebuild)
