import io
import tempfile
import zipfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from services.document_parsers import parse
from train_engine import chunk_id, split_documents


class StructuredIngestionTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.metadata = {"source": "sample", "file_name": "sample"}

    def test_spreadsheet_preserves_headers_sheet_and_unicode(self):
        path = self.root / "sample.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Minerals"
        sheet.append(["Mineral", "State"])
        sheet.append(["लिथियम", "Rajasthan"])
        workbook.save(path)
        documents = parse(path, self.metadata, lambda *_: "")
        self.assertEqual(documents[0].metadata["sheet"], "Minerals")
        self.assertIn("Mineral: लिथियम", documents[0].page_content)
        self.assertEqual(split_documents(documents)[0].metadata["language"], "hi")

    def test_docx_preserves_heading_table_and_embedded_image_text(self):
        path = self.root / "sample.docx"
        xml = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
        <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Mine Safety</w:t></w:r></w:p>
        <w:p><w:r><w:t>Check ventilation before work.</w:t></w:r></w:p>
        <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Hazard</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>Dust</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
        </w:body></w:document>'''
        image = Image.new("RGB", (100, 100), "white")
        output = io.BytesIO()
        image.save(output, format="PNG")
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("word/document.xml", xml)
            archive.writestr("word/media/notice.png", output.getvalue())
        documents = parse(path, self.metadata, lambda *_: "Emergency exit")
        self.assertEqual(documents[0].metadata["heading"], "Mine Safety")
        self.assertEqual(documents[1].metadata["element_type"], "table")
        self.assertIn("Hazard | Dust", documents[1].page_content)
        self.assertEqual(documents[2].metadata["element_type"], "image_ocr")

    def test_opendocument_sheet_and_csv_columns(self):
        ods = self.root / "sample.ods"
        xml = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
        xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
        xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
        <office:body><office:spreadsheet><table:table table:name="Minerals">
        <table:table-row><table:table-cell><text:p>Ore</text:p></table:table-cell></table:table-row>
        <table:table-row><table:table-cell><text:p>Graphite</text:p></table:table-cell></table:table-row>
        </table:table></office:spreadsheet></office:body></office:document-content>'''
        with zipfile.ZipFile(ods, "w") as archive:
            archive.writestr("content.xml", xml)
        self.assertEqual(parse(ods, self.metadata, lambda *_: "")[0].metadata["sheet"], "Minerals")
        csv_file = self.root / "sample.csv"
        csv_file.write_text("Mineral,State\nGraphite,Odisha\n", encoding="utf-8")
        self.assertIn("Mineral: Graphite", parse(csv_file, self.metadata, lambda *_: "")[0].page_content)

    def test_mixed_language_chunks_keep_their_own_language(self):
        path = self.root / "sample.md"
        path.write_text("# English\nThe mine safety system checks the equipment before the shift.\n\n"
                        "# हिंदी\nखनन सुरक्षा के नियम श्रमिकों की रक्षा करते हैं।", encoding="utf-8")
        chunks = split_documents(parse(path, self.metadata, lambda *_: ""))
        self.assertIn("en", [chunk.metadata["language"] for chunk in chunks])
        self.assertIn("hi", [chunk.metadata["language"] for chunk in chunks])

    def test_slide_keeps_title_table_and_image_ocr(self):
        path = self.root / "sample.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Safety"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
        table.cell(0, 0).text = "Hazard"
        table.cell(1, 0).text = "Dust"
        image = Image.new("RGB", (100, 100), "white")
        output = io.BytesIO()
        image.save(output, format="PNG")
        output.seek(0)
        slide.shapes.add_picture(output, Inches(1), Inches(3), Inches(2), Inches(2))
        presentation.save(path)
        documents = parse(path, self.metadata, lambda *_: "Emergency exit")
        self.assertEqual(documents[0].metadata["slide"], 1)
        self.assertIn("Safety", documents[0].page_content)
        self.assertIn("Hazard |", documents[0].page_content)
        self.assertIn("Emergency exit", documents[0].page_content)

    def test_json_paths_and_code_symbols_survive_chunking(self):
        data = self.root / "sample.json"
        data.write_text('{"mine":{"version":"2024","limit":"100 kg"}}', encoding="utf-8")
        documents = parse(data, self.metadata, lambda *_: "")
        self.assertIn("mine.version: 2024", documents[0].page_content)
        code = self.root / "sample.py"
        code.write_text("import math\n\ndef area(radius):\n    return math.pi * radius ** 2\n", encoding="utf-8")
        code_docs = parse(code, self.metadata, lambda *_: "")
        self.assertEqual(code_docs[0].metadata["symbol"], "area")
        self.assertIn("import math", code_docs[0].page_content)
        self.assertEqual(chunk_id(split_documents(code_docs)[0]), chunk_id(split_documents(code_docs)[0]))

    def test_unsupported_format_has_clear_error(self):
        path = self.root / "sample.bin"
        path.write_bytes(b"some bytes")
        with self.assertRaisesRegex(ValueError, "Unsupported file format"):
            parse(path, self.metadata, lambda *_: "")

    def test_scanned_pdf_uses_ocr_and_retains_page(self):
        import fitz
        from train_engine import load_pdf
        path = self.root / "scan.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(path)
        pdf.close()
        with patch("train_engine.ocr_engine_available", return_value=True), patch("train_engine.ocr_image", return_value="खनिज सुरक्षा निर्देश"):
            documents = load_pdf(path, self.root)
        self.assertEqual(documents[0].metadata["page"], 0)
        self.assertTrue(documents[0].metadata["ocr"])
        self.assertIn("खनिज", documents[0].page_content)


if __name__ == "__main__":
    unittest.main()
