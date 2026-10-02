"""Offline, structure-preserving parsers for non-PDF document formats."""

import ast
import csv
import io
import json
import re
import shutil
import subprocess
import tempfile
import tomllib
import xml.etree.ElementTree as ET
import zipfile
import config
import hashlib
from html.parser import HTMLParser
from pathlib import Path

from langchain_core.documents import Document



PARSERS = {}
TEXT_TYPES = {
    ".txt", ".md", ".markdown", ".rst", ".adoc", ".log", ".sql", ".py",
    ".java", ".js", ".jsx", ".ts", ".tsx", ".css", ".c", ".cpp", ".h",
    ".hpp", ".go", ".rs", ".sh", ".ps1", ".bat",
}


def register(*extensions):
    def decorator(function):
        for extension in extensions:
            PARSERS[extension] = function
        return function
    return decorator


def _doc(text, metadata, **details):
    text = str(text).strip()
    if not text:
        return None
    return Document(page_content=text, metadata={**metadata, **details})


def _read(path):
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            pass
    return path.read_text(encoding="utf-8", errors="replace")


def _image_documents(zipped, prefix, metadata, ocr, existing_text=""):
    from PIL import Image
    documents = []
    seen = set()
    seen_images = set()
    for name in zipped.namelist():
        if not name.startswith(prefix) or Path(name).suffix.lower() not in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}:
            continue
        blob = zipped.read(name)
        digest = hashlib.sha1(blob).digest()
        if digest in seen_images:
            continue
        seen_images.add(digest)
        with Image.open(io.BytesIO(blob)) as image:
            text = ocr(image, f"{metadata['source']} image {name}")
        normalized = " ".join(text.casefold().split())
        existing = " ".join(" ".join(doc.page_content.casefold().split()) for doc in documents) + " " + " ".join(existing_text.casefold().split())
        if normalized and normalized not in seen and normalized not in existing:
            seen.add(normalized)
            doc = _doc(text, metadata, element_type="image_ocr", image_name=name, ocr=True)
            if doc:
                documents.append(doc)
    return documents


@register(".docx")
def parse_docx(path, metadata, ocr):
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    documents = []
    section = []
    heading = ""
    def flush():
        if section:
            doc = _doc("\n".join(([heading] if heading else []) + section), metadata, element_type="section", heading=heading)
            if doc:
                documents.append(doc)
            section.clear()
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        body = root.find(f"{namespace}body")
        for child in body:
            if child.tag == f"{namespace}p":
                text = "".join(element.text or "" for element in child.iter(f"{namespace}t")).strip()
                style = child.find(f"{namespace}pPr/{namespace}pStyle")
                style_name = style.attrib.get(f"{namespace}val", "") if style is not None else ""
                if style_name.lower().startswith(("heading", "title")):
                    flush()
                    heading = text
                elif text:
                    if section and sum(len(part) + 1 for part in section) + len(text) > int(config.CHUNK_SIZE):
                        flush()
                    section.append(text)
            elif child.tag == f"{namespace}tbl":
                flush()
                rows = []
                for row in child.findall(f"{namespace}tr"):
                    cells = [" ".join((node.text or "") for node in cell.iter(f"{namespace}t")).strip()
                             for cell in row.findall(f"{namespace}tc")]
                    rows.append(" | ".join(cells))
                doc = _doc("\n".join(([heading] if heading else []) + rows), metadata,
                           element_type="table", heading=heading, table_index=len(documents))
                if doc:
                    documents.append(doc)
        flush()
        documents.extend(_image_documents(archive, "word/media/", metadata, ocr,
                                          "\n".join(doc.page_content for doc in documents)))
    return documents


@register(".xlsx", ".xlsm")
def parse_xlsx(path, metadata, ocr):
    from openpyxl import load_workbook
    documents = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            columns = [str(value).strip() if value is not None else f"Column {index + 1}"
                       for index, value in enumerate(header)]
            batch = []
            start = 2
            def flush(end):
                if not batch:
                    return
                doc = _doc(f"Sheet: {sheet.title}\nColumns: {' | '.join(columns)}\n" + "\n".join(batch),
                           metadata, element_type="table", sheet=sheet.title, row_start=start, row_end=end)
                if doc:
                    documents.append(doc)
                batch.clear()
            for row_number, row in enumerate(rows, start=2):
                values = [f"{columns[index]}: {value}" for index, value in enumerate(row)
                          if value is not None and str(value).strip()]
                if not values:
                    continue
                if not batch:
                    start = row_number
                batch.append("; ".join(values))
                if len(batch) >= 20:
                    flush(row_number)
            flush(sheet.max_row or start)
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        documents.extend(_image_documents(archive, "xl/media/", metadata, ocr,
                                          "\n".join(doc.page_content for doc in documents)))
    return documents


@register(".csv", ".tsv")
def parse_csv(path, metadata, ocr):
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    documents = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter)
        header = next(reader, None)
        if header is None:
            return []
        columns = [name.strip() or f"Column {index + 1}" for index, name in enumerate(header)]
        batch = []
        start = 2
        for row_number, row in enumerate(reader, start=2):
            if not batch:
                start = row_number
            batch.append("; ".join(f"{columns[i] if i < len(columns) else f'Column {i + 1}'}: {value}"
                                   for i, value in enumerate(row) if value.strip()))
            if len(batch) >= 20:
                doc = _doc(f"Columns: {' | '.join(columns)}\n" + "\n".join(batch), metadata,
                           element_type="table", row_start=start, row_end=row_number)
                if doc:
                    documents.append(doc)
                batch.clear()
        if batch:
            doc = _doc(f"Columns: {' | '.join(columns)}\n" + "\n".join(batch), metadata,
                       element_type="table", row_start=start, row_end=row_number)
            if doc:
                documents.append(doc)
    return documents


@register(".pptx")
def parse_pptx(path, metadata, ocr):
    from PIL import Image
    from pptx import Presentation
    documents = []
    seen_images = set()
    for slide_number, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        parts = [f"Slide {slide_number}: {title}"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip() and shape.text.strip() != title:
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if shape.shape_type == 13:
                blob = shape.image.blob
                digest = hashlib.sha1(blob).digest()
                if digest in seen_images:
                    continue
                seen_images.add(digest)
                with Image.open(io.BytesIO(blob)) as image:
                    image_text = ocr(image, f"{metadata['source']} slide {slide_number}")
                if image_text.strip() and image_text.strip() not in parts:
                    parts.append("Image text: " + image_text.strip())
        doc = _doc("\n".join(parts), metadata, element_type="slide", slide=slide_number, slide_title=title)
        if doc:
            documents.append(doc)
    return documents


def _flatten(data, prefix=""):
    if isinstance(data, dict):
        for key, value in data.items():
            yield from _flatten(value, f"{prefix}.{key}" if prefix else str(key))
    elif isinstance(data, list):
        for index, value in enumerate(data):
            yield from _flatten(value, f"{prefix}[{index}]")
    else:
        yield f"{prefix}: {data}"


@register(".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".ipynb")
def parse_structured(path, metadata, ocr):
    suffix = path.suffix.lower()
    content = _read(path)
    if suffix == ".jsonl":
        data = [json.loads(line) for line in content.splitlines() if line.strip()]
    elif suffix in {".json", ".ipynb"}:
        data = json.loads(content)
        if suffix == ".ipynb":
            data = {"cells": [{"cell_type": cell.get("cell_type"), "source": "".join(cell.get("source", []))}
                              for cell in data.get("cells", [])]}
    elif suffix in {".yaml", ".yml"}:
        import yaml
        data = yaml.safe_load(content)
    elif suffix == ".toml":
        data = tomllib.loads(content)
    else:
        import configparser
        parser = configparser.ConfigParser()
        parser.read_string(content)
        data = {section: dict(parser.items(section)) for section in parser.sections()}
    lines = list(_flatten(data))
    return [doc for start in range(0, len(lines), 60)
            if (doc := _doc("\n".join(lines[start:start + 60]), metadata,
                            element_type="hierarchy", path_group=start // 60))]


@register(".xml")
def parse_xml(path, metadata, ocr):
    root = ET.parse(path).getroot()
    lines = []
    def visit(node, parent):
        name = node.tag.split("}")[-1]
        current = f"{parent}/{name}"
        if node.text and node.text.strip():
            lines.append(f"{current}: {node.text.strip()}")
        for key, value in node.attrib.items():
            lines.append(f"{current}/@{key}: {value}")
        for child in node:
            visit(child, current)
    visit(root, "")
    return [doc for start in range(0, len(lines), 60)
            if (doc := _doc("\n".join(lines[start:start + 60]), metadata,
                            element_type="hierarchy", path_group=start // 60))]


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.ignored = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.ignored += 1
        if tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4", "td", "th", "br"}:
            self.parts.append("\n" if tag != "td" else " | ")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.ignored = max(0, self.ignored - 1)

    def handle_data(self, data):
        if not self.ignored:
            self.parts.append(data)


@register(".html", ".htm")
def parse_html(path, metadata, ocr):
    parser = _HTMLText()
    parser.feed(_read(path))
    return [doc] if (doc := _doc("".join(parser.parts), metadata, element_type="html")) else []


@register(".odt", ".ods", ".odp")
def parse_opendocument(path, metadata, ocr):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("content.xml"))
        table_ns = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
        draw_ns = "{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}"
        documents = []
        if path.suffix.lower() == ".ods":
            for sheet in root.iter(f"{table_ns}table"):
                title = sheet.attrib.get(f"{table_ns}name", "Sheet")
                rows = []
                for row in sheet.findall(f"{table_ns}table-row"):
                    cells = [" ".join("".join(cell.itertext()).split())
                             for cell in row.findall(f"{table_ns}table-cell")]
                    if any(cells):
                        rows.append(" | ".join(cells))
                for start in range(0, len(rows), 20):
                    doc = _doc(f"Sheet: {title}\n" + "\n".join(rows[start:start + 20]), metadata,
                               element_type="table", sheet=title, row_start=start + 1,
                               row_end=min(len(rows), start + 20))
                    if doc:
                        documents.append(doc)
        elif path.suffix.lower() == ".odp":
            for slide_number, page in enumerate(root.iter(f"{draw_ns}page"), start=1):
                title = page.attrib.get(f"{draw_ns}name", f"Slide {slide_number}")
                lines = [" ".join("".join(node.itertext()).split()) for node in page.iter()
                         if node.tag.split("}")[-1] in {"p", "h"}]
                doc = _doc(f"Slide {slide_number}: {title}\n" + "\n".join(line for line in lines if line),
                           metadata, element_type="slide", slide=slide_number, slide_title=title)
                if doc:
                    documents.append(doc)
        else:
            lines = []
            for node in root.iter():
                kind = node.tag.split("}")[-1]
                if kind in {"h", "p"}:
                    text = " ".join("".join(node.itertext()).split())
                    if text:
                        lines.append((kind, text))
            documents = [doc for start in range(0, len(lines), 40)
                         if (doc := _doc("\n".join(f"{kind}: {text}" for kind, text in lines[start:start + 40]),
                                         metadata, element_type="section", section_index=start // 40))]
        documents.extend(_image_documents(archive, "Pictures/", metadata, ocr,
                                          "\n".join(doc.page_content for doc in documents)))
        return documents


@register(".rtf")
def parse_rtf(path, metadata, ocr):
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as exc:
        raise RuntimeError("RTF support needs the offline striprtf package installed during setup.") from exc
    return [doc] if (doc := _doc(rtf_to_text(_read(path)), metadata, element_type="prose")) else []


@register(".doc", ".xls", ".ppt")
def parse_legacy_office(path, metadata, ocr):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError(f"{path.suffix.upper()} extraction requires a local LibreOffice installation.")
    target = {".doc": "docx", ".xls": "xlsx", ".ppt": "pptx"}[path.suffix.lower()]
    with tempfile.TemporaryDirectory(prefix="texmin-convert-") as folder:
        profile = (Path(folder) / "profile").as_uri()
        result = subprocess.run([executable, "-env:UserInstallation=" + profile,
                                 "--headless", "--convert-to", target, "--outdir", folder, str(path)],
                                capture_output=True, text=True, timeout=120,
                                creationflags=subprocess.CREATE_NO_WINDOW if __import__("os").name == "nt" else 0)
        converted = Path(folder) / f"{path.stem}.{target}"
        if result.returncode or not converted.is_file():
            raise RuntimeError(f"LibreOffice could not convert {path.name}: {result.stderr[-500:]}")
        return PARSERS[f".{target}"](converted, metadata, ocr)


@register(*TEXT_TYPES)
def parse_text(path, metadata, ocr):
    content = _read(path)
    if path.suffix.lower() == ".py":
        try:
            tree = ast.parse(content)
            lines = content.splitlines()
            imports = [lines[node.lineno - 1] for node in tree.body if isinstance(node, (ast.Import, ast.ImportFrom))]
            documents = []
            for node in tree.body:
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    source = "\n".join(lines[node.lineno - 1:node.end_lineno])
                    doc = _doc("\n".join(imports[:20] + [source]), metadata, element_type="code",
                               symbol=node.name, line_start=node.lineno, line_end=node.end_lineno)
                    if doc:
                        documents.append(doc)
            if documents:
                return documents
        except SyntaxError:
            pass
    if path.suffix.lower() in {".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp", ".sql", ".go", ".rs"}:
        lines = content.splitlines()
        symbols = []
        symbol_pattern = re.compile(r"^\s*(?:export\s+)?(?:async\s+)?(?:class|interface|function|struct|enum|CREATE\s+(?:TABLE|FUNCTION|PROCEDURE))\s+([\w.]+)", re.IGNORECASE)
        for number, line in enumerate(lines, start=1):
            match = symbol_pattern.match(line)
            if match:
                symbols.append((number, match.group(1)))
        if symbols:
            imports = [line for line in lines[:80] if re.match(r"\s*(?:import|from|#include|using)\b", line)]
            documents = []
            for index, (start, symbol) in enumerate(symbols):
                end = symbols[index + 1][0] - 1 if index + 1 < len(symbols) else len(lines)
                doc = _doc("\n".join(imports[:20] + lines[start - 1:end]), metadata,
                           element_type="code", symbol=symbol, line_start=start, line_end=end)
                if doc:
                    documents.append(doc)
            return documents
    sections = re.split(r"(?=^#{1,6}\s|^={3,}\s*$)", content, flags=re.MULTILINE) if path.suffix.lower() in {".md", ".markdown", ".rst"} else [content]
    return [doc for index, section in enumerate(sections)
            if (doc := _doc(section, metadata, element_type="prose" if path.suffix.lower() in {".txt", ".md", ".markdown", ".rst"} else "code",
                            section_index=index))]


def parse(path, metadata, ocr):
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported file format '{path.suffix.lower() or '(none)'}' for {path.name}.")
    return parser(path, metadata, ocr)
